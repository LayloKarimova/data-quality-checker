from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def generate_report(df, errors, output_path="reports/data_quality_report.xlsx"):
    """Создаёт Excel-отчёт по результатам проверки качества данных."""

    # Создаём папку reports, если её нет
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    # Преобразуем список ошибок в DataFrame
    errors_df = pd.DataFrame(errors)

    # Определяем строки Excel, в которых обнаружены ошибки
    error_rows = {error["Row"] for error in errors}

    # Создаём копию исходных данных
    clean_data = df.copy()

    # Добавляем статус каждой строки
    clean_data["Status"] = [
        "Error" if index + 2 in error_rows else "Valid"
        for index in clean_data.index
    ]

    # Общая статистика
    total_rows = len(df)
    rows_with_errors = len(error_rows)
    valid_rows = total_rows - rows_with_errors

    if total_rows > 0:
        quality_score = (valid_rows / total_rows) * 100
    else:
        quality_score = 0

    summary_df = pd.DataFrame({
        "Metric": [
            "Total rows",
            "Valid rows",
            "Rows with errors",
            "Total errors",
            "Data Quality Score"
        ],
        "Value": [
            total_rows,
            valid_rows,
            rows_with_errors,
            len(errors),
            f"{quality_score:.1f}%"
        ]
    })

    # Создаём Excel-файл с тремя листами
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        errors_df.to_excel(
            writer,
            sheet_name="Errors",
            index=False
        )

        clean_data.to_excel(
            writer,
            sheet_name="Clean Data",
            index=False
        )

    # Оформляем созданный Excel
    format_report(output_path)

    return output_path


def format_report(output_path):
    """Добавляет оформление Excel-отчёту."""

    workbook = load_workbook(output_path)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    error_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC"
    )

    valid_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3"
    )

    # Оформляем каждый лист
    for worksheet in workbook.worksheets:

        # Заголовки
        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Закрепляем верхнюю строку
        worksheet.freeze_panes = "A2"

        # Добавляем фильтры
        worksheet.auto_filter.ref = worksheet.dimensions

        # Автоматическая ширина столбцов
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[column_letter].width = (
                min(max_length + 3, 40)
            )

    # Выделяем Status на листе Clean Data
    clean_sheet = workbook["Clean Data"]

    status_column = None

    for cell in clean_sheet[1]:
        if cell.value == "Status":
            status_column = cell.column
            break

    if status_column is not None:
        for row in range(2, clean_sheet.max_row + 1):

            status_cell = clean_sheet.cell(
                row=row,
                column=status_column
            )

            if status_cell.value == "Error":
                status_cell.fill = error_fill
                status_cell.font = Font(bold=True)

            elif status_cell.value == "Valid":
                status_cell.fill = valid_fill

    workbook.save(output_path)